import os
import json
import tempfile

from io import BytesIO

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    send_file,
)

from flask_session import Session

from werkzeug.security import (
    generate_password_hash,
    check_password_hash,
)

from werkzeug.utils import secure_filename

from sqlalchemy.orm import Session as DatabaseSession

from dotenv import load_dotenv

from PyPDF2 import PdfReader

from docx import Document

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from db import engine

from models import (
    Base,
    User,
    AnalysisReport,
)

from ai import (
    analyze_resume,
    validate_resume_document,
    rank_candidates,
    generate_mock_interview,
    evaluate_mock_interview,
    build_ats_resume,
)


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)


app.secret_key = os.getenv(
    "FLASK_SECRET_KEY",
    "change-this-secret-key-before-production",
)


app.config["MAX_CONTENT_LENGTH"] = (
    10 * 1024 * 1024
)


# ============================================================
# SERVER-SIDE SESSION
# ============================================================

SESSION_FOLDER = os.path.join(
    tempfile.gettempdir(),
    "ai_career_copilot_sessions",
)

os.makedirs(
    SESSION_FOLDER,
    exist_ok=True,
)


app.config["SESSION_TYPE"] = "filesystem"

app.config["SESSION_FILE_DIR"] = (
    SESSION_FOLDER
)

app.config["SESSION_PERMANENT"] = False

app.config["SESSION_USE_SIGNER"] = True


Session(app)


# ============================================================
# DATABASE
# ============================================================

Base.metadata.create_all(
    bind=engine
)


def get_db():

    return DatabaseSession(
        bind=engine
    )


# ============================================================
# FILE CONFIGURATION
# ============================================================

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


def allowed_file(filename):

    if not filename:
        return False

    extension = os.path.splitext(
        filename.lower()
    )[1]

    return (
        extension
        in ALLOWED_EXTENSIONS
    )


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_storage):

    try:

        file_storage.stream.seek(0)

        reader = PdfReader(
            file_storage.stream
        )

        text_parts = []

        for page in reader.pages:

            page_text = (
                page.extract_text()
                or ""
            )

            if page_text.strip():

                text_parts.append(
                    page_text
                )

        file_storage.stream.seek(0)

        return "\n".join(
            text_parts
        ).strip()

    except Exception as exc:

        try:
            file_storage.stream.seek(0)
        except Exception:
            pass

        raise ValueError(
            f"Unable to read PDF: {exc}"
        )


def extract_docx_text(file_storage):

    try:

        file_storage.stream.seek(0)

        document = Document(
            file_storage.stream
        )

        text_parts = []

        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                or ""
            ).strip()

            if text:
                text_parts.append(text)

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    value = (
                        cell.text
                        or ""
                    ).strip()

                    if value:
                        row_text.append(
                            value
                        )

                if row_text:

                    text_parts.append(
                        " | ".join(
                            row_text
                        )
                    )

        file_storage.stream.seek(0)

        return "\n".join(
            text_parts
        ).strip()

    except Exception as exc:

        try:
            file_storage.stream.seek(0)
        except Exception:
            pass

        raise ValueError(
            f"Unable to read DOCX: {exc}"
        )


def extract_resume_text(
    file_storage
):

    if not file_storage:
        return ""

    filename = (
        file_storage.filename
        or ""
    )

    extension = os.path.splitext(
        filename.lower()
    )[1]

    if extension == ".pdf":

        return extract_pdf_text(
            file_storage
        )

    if extension == ".docx":

        return extract_docx_text(
            file_storage
        )

    raise ValueError(
        "Only PDF and DOCX Resume files are supported."
    )


# ============================================================
# RESUME VALIDATION
# ============================================================

def validate_resume_or_raise(
    resume_text
):

    text = (
        resume_text
        or ""
    ).strip()

    if not text:

        raise ValueError(
            "No readable Resume text was found."
        )

    if len(text) < 120:

        raise ValueError(
            "The uploaded document contains too little readable text to be treated as a Resume/CV."
        )

    validation = (
        validate_resume_document(
            text
        )
    )

    if not isinstance(
        validation,
        dict,
    ):

        raise ValueError(
            "Resume validation failed. Please try again."
        )

    is_resume = validation.get(
        "is_resume",
        False,
    )

    confidence = validation.get(
        "confidence",
        0,
    )

    try:

        confidence = int(
            confidence
        )

    except Exception:

        confidence = 0

    if (
        not is_resume
        or confidence < 70
    ):

        reason = validation.get(
            "reason",
            "The uploaded document does not appear to be a genuine Resume/CV.",
        )

        raise ValueError(
            reason
        )

    return validation


# ============================================================
# REPORT DATABASE
# ============================================================

def save_analysis_report(
    user_id,
    role,
    result,
    resume_text="",
    filename="",
):

    db = get_db()

    try:

        report = AnalysisReport(

            user_id=user_id,

            target_role=(
                role
                or "General"
            ),

            filename=(
                filename
                or ""
            ),

            resume_text=(
                resume_text
                or ""
            ),

            result=json.dumps(
                result,
                ensure_ascii=False,
            ),
        )

        db.add(report)

        db.commit()

        db.refresh(report)

        return report.id

    finally:

        db.close()


def load_result_json(
    raw_result
):

    if isinstance(
        raw_result,
        dict,
    ):
        return raw_result

    try:

        return json.loads(
            raw_result
            or "{}"
        )

    except Exception:

        return {}


# ============================================================
# PDF HELPERS
# ============================================================

def pdf_escape(value):

    text = str(
        value
        if value is not None
        else ""
    )

    text = text.replace(
        "&",
        "&amp;",
    )

    text = text.replace(
        "<",
        "&lt;",
    )

    text = text.replace(
        ">",
        "&gt;",
    )

    return text


def pdf_list(
    story,
    title,
    items,
    styles,
):

    if not items:
        return

    story.append(
        Paragraph(
            pdf_escape(title),
            styles["Heading2"],
        )
    )

    story.append(
        Spacer(
            1,
            3 * mm,
        )
    )

    if isinstance(
        items,
        str,
    ):

        items = [items]

    for item in items:

        if isinstance(
            item,
            dict,
        ):

            value = " - ".join(
                str(value)
                for value
                in item.values()
                if value
            )

        else:

            value = str(item)

        story.append(
            Paragraph(
                "• "
                + pdf_escape(
                    value
                ),
                styles["BodyText"],
            )
        )

        story.append(
            Spacer(
                1,
                1.5 * mm,
            )
        )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )


def build_analysis_pdf(
    report,
    result,
):

    buffer = BytesIO()

    document = (
        SimpleDocTemplate(
            buffer,

            pagesize=A4,

            rightMargin=18 * mm,

            leftMargin=18 * mm,

            topMargin=18 * mm,

            bottomMargin=18 * mm,
        )
    )

    styles = (
        getSampleStyleSheet()
    )

    story = []


    story.append(
        Paragraph(
            "CareerMind AI",
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            "Resume Analysis Report",
            styles["Heading2"],
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )


    metadata = [

        [
            "Target Role",
            report.target_role
            or "General",
        ],

        [
            "Resume",
            report.filename
            or "Pasted Resume",
        ],

        [
            "ATS Score",
            f'{result.get("ats_score", 0)}/100',
        ],

        [
            "JD Match",
            f'{result.get("jd_match_score", 0)}/100',
        ],
    ]


    table = Table(
        metadata,

        colWidths=[
            40 * mm,
            115 * mm,
        ],
    )


    table.setStyle(

        TableStyle(

            [

                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#EFF6FF"
                    ),
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        "#0F172A"
                    ),
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),

                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(table)

    story.append(
        Spacer(
            1,
            7 * mm,
        )
    )


    summary = result.get(
        "resume_summary",
        "",
    )

    if summary:

        story.append(
            Paragraph(
                "AI Career Summary",
                styles["Heading2"],
            )
        )

        story.append(
            Paragraph(
                pdf_escape(
                    summary
                ),
                styles["BodyText"],
            )
        )

        story.append(
            Spacer(
                1,
                5 * mm,
            )
        )


    section_scores = result.get(
        "section_scores",
        {},
    )

    if section_scores:

        story.append(
            Paragraph(
                "Section Scores",
                styles["Heading2"],
            )
        )

        score_rows = [
            [
                "Section",
                "Score",
            ]
        ]

        for key, value in (
            section_scores.items()
        ):

            score_rows.append(

                [
                    key.replace(
                        "_",
                        " ",
                    ).title(),

                    str(value),
                ]
            )

        score_table = Table(
            score_rows,

            colWidths=[
                90 * mm,
                40 * mm,
            ],
        )

        score_table.setStyle(

            TableStyle(

                [

                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#DBEAFE"
                        ),
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.HexColor(
                            "#CBD5E1"
                        ),
                    ),

                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        story.append(
            score_table
        )

        story.append(
            Spacer(
                1,
                6 * mm,
            )
        )


    pdf_list(
        story,
        "Skills",
        result.get(
            "skills",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Matched Skills",
        result.get(
            "matched_skills",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Missing Skills",
        result.get(
            "missing_skills",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "ATS Keyword Suggestions",
        result.get(
            "keyword_suggestions",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Strengths",
        result.get(
            "strengths",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Weaknesses",
        result.get(
            "weaknesses",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Improvement Suggestions",
        result.get(
            "improvement_suggestions",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Recommended Job Roles",
        result.get(
            "job_recommendations",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Skill Gap Roadmap",
        result.get(
            "roadmap",
            [],
        ),
        styles,
    )

    pdf_list(
        story,
        "Interview Questions",
        result.get(
            "interview_questions",
            [],
        ),
        styles,
    )


    document.build(
        story
    )

    buffer.seek(0)

    return buffer


# ============================================================
# ATS RESUME PDF
# ============================================================

def build_resume_pdf(
    profile,
    resume,
):

    buffer = BytesIO()

    document = (
        SimpleDocTemplate(
            buffer,

            pagesize=A4,

            rightMargin=18 * mm,

            leftMargin=18 * mm,

            topMargin=15 * mm,

            bottomMargin=15 * mm,
        )
    )

    styles = (
        getSampleStyleSheet()
    )

    story = []


    name = profile.get(
        "name",
        "Candidate",
    )

    story.append(
        Paragraph(
            pdf_escape(
                name
            ),
            styles["Title"],
        )
    )


    headline = resume.get(
        "headline",
        "",
    )

    if headline:

        story.append(
            Paragraph(
                pdf_escape(
                    headline
                ),
                styles["Heading2"],
            )
        )


    contact_values = [

        profile.get(
            "email",
            "",
        ),

        profile.get(
            "phone",
            "",
        ),

        profile.get(
            "linkedin",
            "",
        ),
    ]

    contact = " | ".join(
        value
        for value
        in contact_values
        if value
    )

    if contact:

        story.append(
            Paragraph(
                pdf_escape(
                    contact
                ),
                styles["BodyText"],
            )
        )


    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )


    summary = resume.get(
        "professional_summary",
        "",
    )

    if summary:

        story.append(
            Paragraph(
                "Professional Summary",
                styles["Heading2"],
            )
        )

        story.append(
            Paragraph(
                pdf_escape(
                    summary
                ),
                styles["BodyText"],
            )
        )

        story.append(
            Spacer(
                1,
                4 * mm,
            )
        )


    pdf_list(
        story,
        "Skills",
        resume.get(
            "skills",
            [],
        ),
        styles,
    )


    experience = resume.get(
        "experience",
        [],
    )

    if experience:

        story.append(
            Paragraph(
                "Experience",
                styles["Heading2"],
            )
        )

        for item in experience:

            title = item.get(
                "title",
                "",
            )

            organization = item.get(
                "organization",
                "",
            )

            heading = title

            if organization:

                heading += (
                    " - "
                    + organization
                )

            story.append(
                Paragraph(
                    pdf_escape(
                        heading
                    ),
                    styles["Heading3"],
                )
            )

            details = item.get(
                "details",
                [],
            )

            if isinstance(
                details,
                str,
            ):

                details = [
                    details
                ]

            for detail in details:

                story.append(
                    Paragraph(
                        "• "
                        + pdf_escape(
                            detail
                        ),
                        styles["BodyText"],
                    )
                )

            story.append(
                Spacer(
                    1,
                    3 * mm,
                )
            )


    projects = resume.get(
        "projects",
        [],
    )

    if projects:

        story.append(
            Paragraph(
                "Projects",
                styles["Heading2"],
            )
        )

        for item in projects:

            story.append(
                Paragraph(
                    pdf_escape(
                        item.get(
                            "name",
                            "",
                        )
                    ),
                    styles["Heading3"],
                )
            )

            details = item.get(
                "details",
                [],
            )

            if isinstance(
                details,
                str,
            ):

                details = [
                    details
                ]

            for detail in details:

                story.append(
                    Paragraph(
                        "• "
                        + pdf_escape(
                            detail
                        ),
                        styles["BodyText"],
                    )
                )

            story.append(
                Spacer(
                    1,
                    3 * mm,
                )
            )


    pdf_list(
        story,
        "Education",
        resume.get(
            "education",
            [],
        ),
        styles,
    )


    pdf_list(
        story,
        "Certifications",
        resume.get(
            "certifications",
            [],
        ),
        styles,
    )


    pdf_list(
        story,
        "Relevant ATS Keywords",
        resume.get(
            "ats_keywords",
            [],
        ),
        styles,
    )


    document.build(
        story
    )

    buffer.seek(0)

    return buffer


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    if session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "dashboard"
            )
        )

    return redirect(
        url_for(
            "login"
        )
    )


# ============================================================
# SIGNUP
# ============================================================

@app.route(
    "/signup",
    methods=[
        "GET",
        "POST",
    ],
)
def signup():

    if session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "dashboard"
            )
        )


    if request.method == "POST":

        name = (
            request.form.get(
                "name",
                "",
            )
            .strip()
        )

        email = (
            request.form.get(
                "email",
                "",
            )
            .strip()
            .lower()
        )

        password = (
            request.form.get(
                "password",
                "",
            )
        )


        if not name:

            flash(
                "Please enter your full name.",
                "error",
            )

            return render_template(
                "signup.html"
            )


        if not email:

            flash(
                "Please enter your email address.",
                "error",
            )

            return render_template(
                "signup.html"
            )


        if len(password) < 8:

            flash(
                "Password must contain at least 8 characters.",
                "error",
            )

            return render_template(
                "signup.html"
            )


        db = get_db()

        try:

            existing = (

                db.query(User)

                .filter(
                    User.email
                    == email
                )

                .first()
            )


            if existing:

                flash(
                    "An account with this email already exists.",
                    "error",
                )

                return render_template(
                    "signup.html"
                )


            user = User(

                name=name,

                email=email,

                password=(
                    generate_password_hash(
                        password
                    )
                ),
            )


            db.add(user)

            db.commit()

            db.refresh(user)


            session.clear()

            session[
                "user_id"
            ] = user.id

            session[
                "user_name"
            ] = user.name


            flash(
                "Account created successfully.",
                "success",
            )


            return redirect(
                url_for(
                    "dashboard"
                )
            )

        finally:

            db.close()


    return render_template(
        "signup.html"
    )


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=[
        "GET",
        "POST",
    ],
)
def login():

    if session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "dashboard"
            )
        )


    if request.method == "POST":

        email = (
            request.form.get(
                "email",
                "",
            )
            .strip()
            .lower()
        )

        password = request.form.get(
            "password",
            "",
        )


        db = get_db()

        try:

            user = (

                db.query(User)

                .filter(
                    User.email
                    == email
                )

                .first()
            )


            if not user:

                flash(
                    "Invalid email or password.",
                    "error",
                )

                return render_template(
                    "login.html"
                )


            valid_password = False


            try:

                valid_password = (
                    check_password_hash(
                        user.password,
                        password,
                    )
                )

            except Exception:

                valid_password = False


            # Legacy plain text password support
            if (
                not valid_password
                and user.password
                == password
            ):

                valid_password = True

                user.password = (
                    generate_password_hash(
                        password
                    )
                )

                db.commit()


            if not valid_password:

                flash(
                    "Invalid email or password.",
                    "error",
                )

                return render_template(
                    "login.html"
                )


            session.clear()

            session[
                "user_id"
            ] = user.id

            session[
                "user_name"
            ] = user.name


            flash(
                "Login successful.",
                "success",
            )


            return redirect(
                url_for(
                    "dashboard"
                )
            )

        finally:

            db.close()


    return render_template(
        "login.html"
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route(
    "/logout"
)
def logout():

    session.clear()

    flash(
        "You have been logged out.",
        "success",
    )

    return redirect(
        url_for(
            "login"
        )
    )


# ============================================================
# DASHBOARD
# ============================================================

@app.route(
    "/dashboard",
    methods=[
        "GET",
        "POST",
    ],
)
def dashboard():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    result = None

    role = ""

    job_description = ""


    if request.method == "POST":

        role = (
            request.form.get(
                "role",
                "",
            )
            .strip()
        )

        job_description = (
            request.form.get(
                "job_description",
                "",
            )
            .strip()
        )

        pasted_text = (
            request.form.get(
                "text",
                "",
            )
            .strip()
        )

        uploaded_file = (
            request.files.get(
                "file"
            )
        )

        filename = ""


        try:

            resume_text = ""


            if (
                uploaded_file
                and uploaded_file.filename
            ):

                filename = (
                    secure_filename(
                        uploaded_file.filename
                    )
                )

                if not allowed_file(
                    filename
                ):

                    raise ValueError(
                        "Only PDF and DOCX Resume files are supported."
                    )

                resume_text = (
                    extract_resume_text(
                        uploaded_file
                    )
                )

            elif pasted_text:

                resume_text = (
                    pasted_text
                )

            else:

                raise ValueError(
                    "Please upload a Resume/CV or paste Resume text."
                )


            validate_resume_or_raise(
                resume_text
            )


            result = analyze_resume(

                resume_text,

                role,

                job_description,
            )


            if not isinstance(
                result,
                dict,
            ):

                result = {
                    "error":
                    "AI analysis returned an invalid response."
                }


            if not result.get(
                "error"
            ):

                report_id = (
                    save_analysis_report(

                        session[
                            "user_id"
                        ],

                        role,

                        result,

                        resume_text,

                        filename,
                    )
                )


                result[
                    "report_id"
                ] = report_id


                # Server-side session allows large Resume content safely.
                session[
                    "latest_resume_text"
                ] = resume_text

                session[
                    "latest_target_role"
                ] = role


        except Exception as exc:

            result = {
                "error":
                str(exc)
            }


    return render_template(

        "dashboard.html",

        result=result,

        role=role,

        job_description=(
            job_description
        ),
    )


# ============================================================
# HISTORY
# ============================================================

@app.route(
    "/history"
)
def history():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    db = get_db()

    try:

        records = (

            db.query(
                AnalysisReport
            )

            .filter(
                AnalysisReport.user_id
                == session[
                    "user_id"
                ]
            )

            .order_by(
                AnalysisReport.id.desc()
            )

            .all()
        )


        reports = []


        for report in records:

            reports.append(

                {

                    "id":
                    report.id,

                    "target_role":
                    report.target_role,

                    "filename":
                    report.filename,

                    "result":
                    load_result_json(
                        report.result
                    ),
                }
            )


    finally:

        db.close()


    return render_template(

        "history.html",

        reports=reports,
    )


# ============================================================
# DOWNLOAD ANALYSIS PDF
# ============================================================

@app.route(
    "/report/<int:report_id>/pdf"
)
def download_report_pdf(
    report_id
):

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    db = get_db()

    try:

        report = (

            db.query(
                AnalysisReport
            )

            .filter(
                AnalysisReport.id
                == report_id,

                AnalysisReport.user_id
                == session[
                    "user_id"
                ],
            )

            .first()
        )


        if not report:

            return (
                "Report not found.",
                404,
            )


        result = (
            load_result_json(
                report.result
            )
        )


        pdf_buffer = (
            build_analysis_pdf(
                report,
                result,
            )
        )


        filename = (
            f"resume_analysis_{report.id}.pdf"
        )


        return send_file(

            pdf_buffer,

            as_attachment=True,

            download_name=filename,

            mimetype=(
                "application/pdf"
            ),
        )

    finally:

        db.close()


# ============================================================
# RECRUITER - MULTIPLE RESUME RANKING
# ============================================================

@app.route(
    "/recruiter",
    methods=[
        "GET",
        "POST",
    ],
)
def recruiter():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    result = None

    rejected = []

    role = ""

    job_description = ""


    if request.method == "POST":

        role = (
            request.form.get(
                "role",
                "",
            )
            .strip()
        )

        job_description = (
            request.form.get(
                "job_description",
                "",
            )
            .strip()
        )


        uploaded_files = (
            request.files.getlist(
                "resumes"
            )
        )


        uploaded_files = [

            item

            for item
            in uploaded_files

            if (
                item
                and item.filename
            )
        ]


        if not role:

            result = {
                "error":
                "Please enter a target job role."
            }


        elif len(
            uploaded_files
        ) < 2:

            result = {
                "error":
                "Please select at least two Resume/CV files."
            }


        elif len(
            uploaded_files
        ) > 10:

            result = {
                "error":
                "A maximum of 10 Resume/CV files can be analyzed at one time."
            }


        else:

            candidates = []


            for index, file in enumerate(
                uploaded_files,
                start=1,
            ):

                original_filename = (
                    file.filename
                    or f"candidate_{index}"
                )

                safe_filename = (
                    secure_filename(
                        original_filename
                    )
                )


                if not allowed_file(
                    safe_filename
                ):

                    rejected.append(

                        {
                            "filename":
                            original_filename,

                            "reason":
                            "Unsupported file type. Only PDF and DOCX are allowed.",
                        }
                    )

                    continue


                try:

                    resume_text = (
                        extract_resume_text(
                            file
                        )
                    )


                    validate_resume_or_raise(
                        resume_text
                    )


                    candidates.append(

                        {
                            "candidate_id":
                            index,

                            "filename":
                            original_filename,

                            "resume_text":
                            resume_text,
                        }
                    )


                except Exception as exc:

                    rejected.append(

                        {
                            "filename":
                            original_filename,

                            "reason":
                            str(exc),
                        }
                    )


            if len(
                candidates
            ) < 2:

                result = {

                    "error":
                    "At least two valid Resume/CV files are required."
                }


            else:

                try:

                    result = rank_candidates(

                        candidates,

                        role,

                        job_description,
                    )


                    if not isinstance(
                        result,
                        dict,
                    ):

                        result = {
                            "error":
                            "Candidate ranking returned an invalid response."
                        }


                except Exception as exc:

                    result = {
                        "error":
                        str(exc)
                    }


    return render_template(

        "recruiter.html",

        result=result,

        rejected=rejected,

        role=role,

        job_description=(
            job_description
        ),
    )


# ============================================================
# ORAL MOCK INTERVIEW
# ============================================================

@app.route(
    "/mock-interview",
    methods=[
        "GET",
        "POST",
    ],
)
def mock_interview():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    questions = None

    evaluation = None

    role = (
        session.get(
            "mock_role",
            "",
        )
    )


    if request.method == "POST":

        action = (
            request.form.get(
                "action",
                "generate",
            )
        )


        # ====================================================
        # GENERATE INTERVIEW QUESTIONS
        # ====================================================

        if action == "generate":

            role = (
                request.form.get(
                    "role",
                    "",
                )
                .strip()
            )


            uploaded_file = (
                request.files.get(
                    "file"
                )
            )


            pasted_text = (
                request.form.get(
                    "text",
                    "",
                )
                .strip()
            )


            try:

                resume_text = ""


                if (
                    uploaded_file
                    and uploaded_file.filename
                ):

                    filename = (
                        secure_filename(
                            uploaded_file.filename
                        )
                    )


                    if not allowed_file(
                        filename
                    ):

                        raise ValueError(
                            "Only PDF and DOCX Resume files are supported."
                        )


                    resume_text = (
                        extract_resume_text(
                            uploaded_file
                        )
                    )


                elif pasted_text:

                    resume_text = (
                        pasted_text
                    )


                else:

                    resume_text = (
                        session.get(
                            "latest_resume_text",
                            "",
                        )
                    )


                if not resume_text:

                    raise ValueError(
                        "Please upload a Resume/CV, paste Resume text, or analyze a Resume from the Dashboard first."
                    )


                validate_resume_or_raise(
                    resume_text
                )


                generated = (
                    generate_mock_interview(

                        resume_text,

                        role,
                    )
                )


                if not isinstance(
                    generated,
                    dict,
                ):

                    raise ValueError(
                        "AI returned an invalid interview response."
                    )


                questions = (
                    generated.get(
                        "questions",
                        [],
                    )
                )


                if not questions:

                    raise ValueError(
                        "No interview questions were generated."
                    )


                # Store in server-side session for final evaluation.
                session[
                    "mock_resume_text"
                ] = resume_text

                session[
                    "mock_role"
                ] = role

                session[
                    "mock_questions"
                ] = questions


            except Exception as exc:

                flash(
                    str(exc),
                    "error",
                )

                questions = None


        # ====================================================
        # EVALUATE ORAL ANSWERS
        # ====================================================

        elif action == "evaluate":

            resume_text = (
                session.get(
                    "mock_resume_text",
                    "",
                )
            )

            role = (
                session.get(
                    "mock_role",
                    "",
                )
            )

            questions = (
                session.get(
                    "mock_questions",
                    [],
                )
            )


            if (
                not resume_text
                or not questions
            ):

                flash(
                    "Interview session expired. Please start a new interview.",
                    "error",
                )

                return redirect(
                    url_for(
                        "mock_interview"
                    )
                )


            answers = []


            for index, question in enumerate(
                questions,
                start=1,
            ):

                question_id = (
                    question.get(
                        "id",
                        index,
                    )
                )


                answer = (
                    request.form.get(
                        f"answer_{question_id}",
                        "",
                    )
                    .strip()
                )


                answers.append(

                    {
                        "id":
                        question_id,

                        "question":
                        question.get(
                            "question",
                            "",
                        ),

                        "answer":
                        answer,
                    }
                )


            answered_count = sum(

                1

                for item
                in answers

                if item[
                    "answer"
                ]
            )


            if answered_count == 0:

                flash(
                    "No interview answers were received.",
                    "error",
                )

                return render_template(

                    "mock_interview.html",

                    questions=questions,

                    role=role,

                    evaluation=None,
                )


            try:

                evaluation = (
                    evaluate_mock_interview(

                        resume_text,

                        role,

                        questions,

                        answers,
                    )
                )


                if not isinstance(
                    evaluation,
                    dict,
                ):

                    raise ValueError(
                        "AI returned an invalid interview evaluation."
                    )


                # Interview has completed.
                session.pop(
                    "mock_resume_text",
                    None,
                )

                session.pop(
                    "mock_questions",
                    None,
                )

                session.pop(
                    "mock_role",
                    None,
                )


                # Hide live interview UI
                # after final evaluation.
                questions = None


            except Exception as exc:

                flash(
                    str(exc),
                    "error",
                )


    return render_template(

        "mock_interview.html",

        questions=questions,

        evaluation=evaluation,

        role=role,
    )


# ============================================================
# ATS RESUME BUILDER
# ============================================================

@app.route(
    "/resume-builder",
    methods=[
        "GET",
        "POST",
    ],
)
def resume_builder():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    result = None

    profile = (
        session.get(
            "resume_builder_profile"
        )
    )


    if request.method == "POST":

        profile = {

            "name":
            request.form.get(
                "name",
                "",
            ).strip(),

            "email":
            request.form.get(
                "email",
                "",
            ).strip(),

            "phone":
            request.form.get(
                "phone",
                "",
            ).strip(),

            "linkedin":
            request.form.get(
                "linkedin",
                "",
            ).strip(),

            "target_role":
            request.form.get(
                "target_role",
                "",
            ).strip(),

            "summary":
            request.form.get(
                "summary",
                "",
            ).strip(),

            "skills":
            request.form.get(
                "skills",
                "",
            ).strip(),

            "experience":
            request.form.get(
                "experience",
                "",
            ).strip(),

            "projects":
            request.form.get(
                "projects",
                "",
            ).strip(),

            "education":
            request.form.get(
                "education",
                "",
            ).strip(),

            "certifications":
            request.form.get(
                "certifications",
                "",
            ).strip(),
        }


        if not profile[
            "name"
        ]:

            result = {
                "error":
                "Please enter your full name."
            }


        elif not profile[
            "target_role"
        ]:

            result = {
                "error":
                "Please enter your target job role."
            }


        else:

            try:

                result = (
                    build_ats_resume(
                        profile
                    )
                )


                if not isinstance(
                    result,
                    dict,
                ):

                    result = {
                        "error":
                        "AI Resume Builder returned an invalid response."
                    }


                if not result.get(
                    "error"
                ):

                    session[
                        "resume_builder_profile"
                    ] = profile

                    session[
                        "resume_builder_result"
                    ] = result


            except Exception as exc:

                result = {
                    "error":
                    str(exc)
                }


    return render_template(

        "resume_builder.html",

        result=result,

        profile=profile,
    )


# ============================================================
# RESUME BUILDER PDF
# ============================================================

@app.route(
    "/resume-builder/pdf"
)
def resume_builder_pdf():

    if not session.get(
        "user_id"
    ):

        return redirect(
            url_for(
                "login"
            )
        )


    profile = (
        session.get(
            "resume_builder_profile"
        )
    )

    resume = (
        session.get(
            "resume_builder_result"
        )
    )


    if not profile or not resume:

        flash(
            "Please build your Resume before downloading the PDF.",
            "error",
        )

        return redirect(
            url_for(
                "resume_builder"
            )
        )


    pdf_buffer = (
        build_resume_pdf(
            profile,
            resume,
        )
    )


    safe_name = (
        secure_filename(
            profile.get(
                "name",
                "resume",
            )
        )
        or "resume"
    )


    return send_file(

        pdf_buffer,

        as_attachment=True,

        download_name=(
            f"{safe_name}_ATS_Resume.pdf"
        ),

        mimetype=(
            "application/pdf"
        ),
    )


# ============================================================
# 413 FILE TOO LARGE
# ============================================================

@app.errorhandler(413)
def file_too_large(
    error
):

    return (
        render_template(
            "base.html"
        ),
        413,
    )


# ============================================================
# 404
# ============================================================

@app.errorhandler(404)
def page_not_found(
    error
):

    return (
        "Page not found.",
        404,
    )


# ============================================================
# START APPLICATION
# ============================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )